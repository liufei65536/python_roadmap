from openpyxl import Workbook


wb = Workbook()
ws = wb.active
treeData = [["Type", "Leaf Color", "Height"], ["Maple", "Red", 549],
            ["Oak", "Green", 783], ["Pine", "Green", 1204]]

# 将数据追加写入wb
for row in treeData:
    ws.append(row)

# 改变字体
from openpyxl.styles import Font
ft = Font(bold=True)
for row in ws["A1:C1"]:
    for cell in row:
        cell.font = ft

# 绘制图标 BarChart()
from openpyxl.chart import BarChart, Series, Reference
chart = BarChart()
chart.type = "col"
chart.title = "Tree Height"
chart.y_axis.title = 'Height (cm)'
chart.x_axis.title = 'Tree Type'
chart.legend = None
data = Reference(ws, min_col=3, min_row=2, max_row=4, max_col=3)
categories = Reference(ws, min_col=1, min_row=2, max_row=4, max_col=1)
chart.add_data(data)
chart.set_categories(categories)
ws.add_chart(chart, "E1")
wb.save("TreeData.xlsx")